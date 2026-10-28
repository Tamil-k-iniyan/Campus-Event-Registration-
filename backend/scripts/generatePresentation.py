import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# 1. Initialize presentation (16:9 Widescreen)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_slide_layout = prs.slide_layouts[6]

# Theme Colors
NAVY = RGBColor(15, 23, 42)        # #0F172A
INDIGO = RGBColor(30, 41, 59)      # #1E293B
BRAND_BLUE = RGBColor(29, 78, 216) # #1D4ED8
BG_LIGHT = RGBColor(248, 250, 252) # #F8FAFC
CARD_BG = RGBColor(255, 255, 255)
CARD_BORDER = RGBColor(226, 232, 240)
TEXT_DARK = RGBColor(15, 23, 42)
TEXT_MUTED = RGBColor(100, 116, 139)
TEXT_LIGHT = RGBColor(255, 255, 255)
GREEN = RGBColor(5, 150, 105)
AMBER = RGBColor(217, 119, 6)

SCREENSHOTS_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '../../docs/screenshots'))

def add_header(slide, title_text, category_text="WEB TECHNOLOGY PROJECT"):
    """Adds consistent professional header to content slides"""
    # Background fill
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_LIGHT
    bg.line.fill.background()

    # Header top bar
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = NAVY
    top_bar.line.fill.background()

    # Category Pill
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.18), Inches(2.8), Inches(0.28))
    pill.fill.solid()
    pill.fill.fore_color.rgb = BRAND_BLUE
    pill.line.fill.background()
    p_tf = pill.text_frame
    p_tf.word_wrap = True
    p = p_tf.paragraphs[0]
    p.text = category_text.upper()
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT
    p.alignment = PP_ALIGN.CENTER

    # Title text
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.48), Inches(11.5), Inches(0.55))
    tf = txBox.text_frame
    tf.word_wrap = True
    p2 = tf.paragraphs[0]
    p2.text = title_text
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_LIGHT

    # Footer line
    ft_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.1), Inches(11.7), Inches(0.3))
    ft_tf = ft_box.text_frame
    p3 = ft_tf.paragraphs[0]
    p3.text = "Campus Event Registration and Management System  |  Karunya University  |  Tamilzhniyan K. (URK24CS1140)"
    p3.font.size = Pt(9.5)
    p3.font.color.rgb = TEXT_MUTED

def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
    """Creates a rounded styled card container"""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    return card

def set_speaker_notes(slide, notes_text):
    """Adds speaker notes to slide"""
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = notes_text

# ==============================================================================
# SLIDE 1: TITLE SLIDE
# ==============================================================================
s1 = prs.slides.add_slide(blank_slide_layout)
# Dark navy background
s1_bg = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
s1_bg.fill.solid()
s1_bg.fill.fore_color.rgb = NAVY
s1_bg.line.fill.background()

# Decorative accent bar
bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.0), Inches(0.15), Inches(4.8))
bar.fill.solid()
bar.fill.fore_color.rgb = BRAND_BLUE
bar.line.fill.background()

# Title text box
s1_title_box = s1.shapes.add_textbox(Inches(1.2), Inches(0.9), Inches(11.0), Inches(2.2))
s1_tf = s1_title_box.text_frame
s1_tf.word_wrap = True

p = s1_tf.paragraphs[0]
p.text = "WEB TECHNOLOGY CAPSTONE PROJECT"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = RGBColor(56, 189, 248) # Light blue accent

p2 = s1_tf.add_paragraph()
p2.text = "Campus Event Registration and\nManagement System"
p2.font.size = Pt(36)
p2.font.bold = True
p2.font.color.rgb = TEXT_LIGHT

p3 = s1_tf.add_paragraph()
p3.text = "A Modern Full-Stack MERN Platform with Concurrency-Safe Atomic Reservations"
p3.font.size = Pt(15)
p3.font.color.rgb = RGBColor(203, 213, 225)

# Details Card
card1 = add_card(s1, Inches(1.2), Inches(3.6), Inches(10.8), Inches(2.4), bg_color=INDIGO, border_color=BRAND_BLUE)
c1_box = s1.shapes.add_textbox(Inches(1.5), Inches(3.75), Inches(10.2), Inches(2.1))
c1_tf = c1_box.text_frame
c1_tf.word_wrap = True

def add_detail_row(tf, label, value):
    p = tf.add_paragraph() if len(tf.paragraphs[0].text) > 0 else tf.paragraphs[0]
    run1 = p.add_run()
    run1.text = f"{label}:  "
    run1.font.bold = True
    run1.font.size = Pt(13)
    run1.font.color.rgb = RGBColor(148, 163, 184)
    run2 = p.add_run()
    run2.text = value
    run2.font.bold = True
    run2.font.size = Pt(13.5)
    run2.font.color.rgb = TEXT_LIGHT

add_detail_row(c1_tf, "Candidate Name", "Tamilzhniyan K.")
add_detail_row(c1_tf, "Register / Roll Number", "URK24CS1140")
add_detail_row(c1_tf, "Department & Branch", "B.Tech in Computer Science and Engineering")
add_detail_row(c1_tf, "Institution", "Karunya University  |  Academic Year 2026")
add_detail_row(c1_tf, "GitHub Repository", "https://github.com/Tamil-k-iniyan/Campus-Event-Registration-")

set_speaker_notes(s1, "Respected examiners, good morning. I am Tamilzhniyan K., register number URK24CS1140 from the Department of Computer Science and Engineering at Karunya University. Today, I present my Web Technology capstone project: the Campus Event Registration and Management System, built with the MERN stack.")

# ==============================================================================
# SLIDE 2: PROJECT OVERVIEW
# ==============================================================================
s2 = prs.slides.add_slide(blank_slide_layout)
add_header(s2, "Project Overview: Unified Campus Event Ecosystem", "Context & Significance")

card_l = add_card(s2, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.2))
box_l = s2.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.0), Inches(4.8))
tf_l = box_l.text_frame
tf_l.word_wrap = True
p = tf_l.paragraphs[0]
p.text = "What is CampusEvents?"
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = BRAND_BLUE

items_l = [
    ("Centralized Digital Portal", "Replaces disorganized WhatsApp posters, physical flyers, and disconnected Google Forms with one official hub."),
    ("Two Segregated Roles", "Role-Based Access Control dividing Student Attendees from Faculty Event Administrators."),
    ("Real-Time Capacity Gauges", "Students immediately see live remaining seats before attempting registration."),
    ("Instant Self-Service", "One-click registration, instant cancellation before start time, and downloadable CSV rosters for faculty.")
]
for title, desc in items_l:
    p = tf_l.add_paragraph()
    p.space_before = Pt(12)
    run = p.add_run()
    run.text = f"• {title}: "
    run.font.bold = True
    run.font.size = Pt(12.5)
    run.font.color.rgb = TEXT_DARK
    run2 = p.add_run()
    run2.text = desc
    run2.font.size = Pt(12)
    run2.font.color.rgb = TEXT_MUTED

card_r = add_card(s2, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.2))
box_r = s2.shapes.add_textbox(Inches(7.1), Inches(1.7), Inches(5.1), Inches(4.8))
tf_r = box_r.text_frame
tf_r.word_wrap = True
p = tf_r.paragraphs[0]
p.text = "System Interaction Flow"
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = GREEN

flow_steps = [
    ("1. Student User", "Discovers hackathons, workshops, and sports meets. Submits booking with optional special notes."),
    ("2. React Client (Vite)", "Dispatches asynchronous requests with JWT Bearer tokens to REST endpoints."),
    ("3. Express.js Server", "Validates inputs, sanitizes data, and executes atomic conditional seat increments."),
    ("4. MongoDB Database", "Stores persistent User profiles, Event catalogs, and Unique Registration records."),
    ("5. Faculty Admin", "Monitors capacity utilization, reviews attendee rosters, and exports CSV reports.")
]
for step, desc in flow_steps:
    p = tf_r.add_paragraph()
    p.space_before = Pt(10)
    run = p.add_run()
    run.text = f"{step}\n"
    run.font.bold = True
    run.font.size = Pt(12.5)
    run.font.color.rgb = BRAND_BLUE
    run2 = p.add_run()
    run2.text = desc
    run2.font.size = Pt(11.5)
    run2.font.color.rgb = TEXT_MUTED

set_speaker_notes(s2, "This slide outlines what the Campus Event Registration and Management System achieves. It replaces paper notices and spreadsheets with a clean digital portal connecting students and faculty. The diagram demonstrates the 5-step lifecycle from student discovery to atomic database writes and administrative CSV exports.")

# ==============================================================================
# SLIDE 3: PROBLEM STATEMENT
# ==============================================================================
s3 = prs.slides.add_slide(blank_slide_layout)
add_header(s3, "Problem Statement: Challenges in Legacy Event Management", "Problem Analysis")

problems = [
    ("Paper & Google Forms Vulnerabilities", "Traditional forms lack atomic capacity validation. When 50 students apply for 30 seats simultaneously, overbooking occurs, causing chaos at the venue.", AMBER),
    ("Duplicate Registration Submissions", "Without database-level unique compound keys, students register multiple times under different names, corrupting attendance records.", RGBColor(225, 29, 72)),
    ("Zero Visibility & Missed Deadlines", "Notices are scattered on bulletin boards and social media groups. Students miss deadlines and remaining seat counts are invisible.", BRAND_BLUE),
    ("Severe Administrative Overhead", "Faculty and student coordinators spend hours manually compiling attendance rosters and validating student roll numbers.", INDIGO)
]

for i, (title, desc, color) in enumerate(problems):
    col = i % 2
    row = i // 2
    left = Inches(0.8 + col * 6.0)
    top = Inches(1.5 + row * 2.7)
    add_card(s3, left, top, Inches(5.7), Inches(2.4))
    
    # Accent indicator
    accent = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.12), Inches(2.4))
    accent.fill.solid()
    accent.fill.fore_color.rgb = color
    accent.line.fill.background()

    tb = s3.shapes.add_textbox(left + Inches(0.3), top + Inches(0.15), Inches(5.2), Inches(2.1))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    p2 = tf.add_paragraph()
    p2.space_before = Pt(6)
    p2.text = desc
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_MUTED

set_speaker_notes(s3, "Here we highlight the core bottlenecks of current campus processes: lack of capacity enforcement causing overbooking, duplicate student submissions, scattered discovery, and manual attendance reconciliation. Our project directly solves each of these vulnerabilities through software engineering.")

# ==============================================================================
# SLIDE 4: PROJECT OBJECTIVES
# ==============================================================================
s4 = prs.slides.add_slide(blank_slide_layout)
add_header(s4, "Project Objectives: Engineering Goals & Deliverables", "Scope & Goals")

objs = [
    ("1. Centralized Event Hub", "Establish a single searchable catalog for academic seminars, hackathons, workshops, cultural galas, and athletic championships."),
    ("2. Concurrency-Safe Booking", "Implement atomic database reservation operators to eliminate overselling under simultaneous peak student registrations."),
    ("3. Duplicate Prevention", "Enforce compound unique database indexing ({ student: 1, event: 1 }) ensuring zero duplicate active registrations."),
    ("4. Role-Based Access Control", "Segregate student attendees from faculty administrators using cryptographically signed JSON Web Tokens (JWT)."),
    ("5. Live Faculty Administration", "Provide real-time analytics dashboards, capacity progress tracking, and 1-click CSV roster generation."),
    ("6. High-Performance Responsive UI", "Develop an intuitive user experience using React 18, Vite, and Bootstrap 5 adaptable to mobile, tablet, and desktop viewports.")
]

for i, (title, desc) in enumerate(objs):
    col = i % 2
    row = i // 2
    left = Inches(0.8 + col * 6.0)
    top = Inches(1.5 + row * 1.8)
    add_card(s4, left, top, Inches(5.7), Inches(1.6))

    tb = s4.shapes.add_textbox(left + Inches(0.3), top + Inches(0.15), Inches(5.2), Inches(1.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE
    
    p2 = tf.add_paragraph()
    p2.space_before = Pt(4)
    p2.text = desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_MUTED

set_speaker_notes(s4, "The project accomplished six key engineering objectives: centralizing event listings, guaranteeing atomic seat capacity reservation, preventing duplicate registrations via unique indexes, implementing stateless JWT authentication, providing faculty administration with CSV export, and delivering a responsive React user interface.")

# ==============================================================================
# SLIDE 5: PROPOSED SYSTEM & KEY FEATURES
# ==============================================================================
s5 = prs.slides.add_slide(blank_slide_layout)
add_header(s5, "Proposed System: Student & Administrator Capabilities", "System Features")

# Student Column
c_s = add_card(s5, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.2))
tb_s = s5.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.1), Inches(4.8))
tf_s = tb_s.text_frame
tf_s.word_wrap = True
p = tf_s.paragraphs[0]
p.text = "Student Portal Features"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = BRAND_BLUE

s_feats = [
    ("Account Registration & Auth", "Secure login with JWT token session handling and profile management."),
    ("Search & Multi-Filter Catalog", "Filter by Technical, Hackathon, Cultural, Sports, and Seminar categories."),
    ("Seat Capacity Indicators", "Color-coded progress bars showing exact seats remaining in real time."),
    ("1-Click Atomic Registration", "Instant seat reservation with optional special dietary or team notes."),
    ("Self-Service Cancellation", "Relinquish seats before event start date, restoring seat availability."),
    ("Student Dashboard", "Personal audit log of confirmed, upcoming, and past event bookings.")
]
for title, desc in s_feats:
    p = tf_s.add_paragraph()
    p.space_before = Pt(6)
    r1 = p.add_run()
    r1.text = f"✔ {title}: "
    r1.font.bold = True
    r1.font.size = Pt(11.5)
    r1.font.color.rgb = TEXT_DARK
    r2 = p.add_run()
    r2.text = desc
    r2.font.size = Pt(11)
    r2.font.color.rgb = TEXT_MUTED

# Admin Column
c_a = add_card(s5, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.2))
tb_a = s5.shapes.add_textbox(Inches(7.1), Inches(1.7), Inches(5.1), Inches(4.8))
tf_a = tb_a.text_frame
tf_a.word_wrap = True
p = tf_a.paragraphs[0]
p.text = "Faculty Admin Console"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = GREEN

a_feats = [
    ("Executive KPI Dashboard", "Instant overview of Total Events, Upcoming, Registered Students, and Bookings."),
    ("Full Event CRUD Management", "Create, edit, publish, and delete events with capacity bound enforcement."),
    ("Attendee Roster Inspection", "Search registered students by full name, college email, or roll number."),
    ("Automated CSV Export", "Single-click download of attendance spreadsheets formatted for Excel."),
    ("Publication Status Toggles", "Switch events between Published, Draft, Closed, and Cancelled modes."),
    ("Safety Bounds Checking", "Prevents reducing event capacity below the number of enrolled attendees.")
]
for title, desc in a_feats:
    p = tf_a.add_paragraph()
    p.space_before = Pt(6)
    r1 = p.add_run()
    r1.text = f"✔ {title}: "
    r1.font.bold = True
    r1.font.size = Pt(11.5)
    r1.font.color.rgb = TEXT_DARK
    r2 = p.add_run()
    r2.text = desc
    r2.font.size = Pt(11)
    r2.font.color.rgb = TEXT_MUTED

set_speaker_notes(s5, "Slide 5 details the functional breakdown. The student portal focuses on frictionless discovery, atomic reservation, and booking history. The faculty administration panel gives organizers complete control: editing event dates, monitoring capacity bars, inspecting student rosters, and downloading CSV reports.")

# ==============================================================================
# SLIDE 6: TECHNOLOGY STACK
# ==============================================================================
s6 = prs.slides.add_slide(blank_slide_layout)
add_header(s6, "Technology Stack: Full-Stack MERN Architecture", "Core Technologies")

techs = [
    ("React 18 & Vite", "Frontend Presentation", "Component-based architecture, virtual DOM for instant UI updates, client-side routing with React Router DOM, bundled via Vite 6 Rollup for sub-2s build times.", BRAND_BLUE),
    ("Node.js & Express.js", "Backend Application Tier", "Asynchronous, non-blocking REST API server. Implements middleware for Helmet, CORS, body parsing, express-rate-limit, and centralized error catching.", GREEN),
    ("MongoDB & Mongoose 8", "Database Persistence", "NoSQL document store with Mongoose ODM schemas. Enforces compound unique indexes, timestamps, virtual computed properties, and atomic update operators.", AMBER),
    ("Security & Tooling", "Authentication & CI/CD", "Bcrypt.js password hashing (10 salt rounds), stateless JSON Web Tokens (7-day expiry), Git version control, and GitHub Actions continuous deployment.", INDIGO)
]

for i, (name, role, desc, color) in enumerate(techs):
    left = Inches(0.8 + i * 2.95)
    top = Inches(1.6)
    add_card(s6, left, top, Inches(2.8), Inches(5.1))

    # Top accent header
    hdr = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(2.8), Inches(0.85))
    hdr.fill.solid()
    hdr.fill.fore_color.rgb = color
    hdr.line.fill.background()
    h_tf = hdr.text_frame
    h_tf.word_wrap = True
    p = h_tf.paragraphs[0]
    p.text = name
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT
    p.alignment = PP_ALIGN.CENTER

    tb = s6.shapes.add_textbox(left + Inches(0.15), top + Inches(0.95), Inches(2.5), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = role.upper()
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = color
    
    p2 = tf.add_paragraph()
    p2.space_before = Pt(8)
    p2.text = desc
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = TEXT_MUTED

set_speaker_notes(s6, "Here is our technical foundation: React 18 with Vite for frontend, Node.js and Express.js for the REST API backend, MongoDB with Mongoose 8 for data modeling, and bcrypt with JWT for security. Each technology was selected for performance, reliability, and academic rigor.")

# ==============================================================================
# SLIDE 7: SYSTEM ARCHITECTURE
# ==============================================================================
s7 = prs.slides.add_slide(blank_slide_layout)
add_header(s7, "System Architecture: 3-Tier Multi-Layer Design", "Architectural Model")

tiers = [
    ("Tier 1: Client Presentation Tier", "React 18 Single Page Application (Port 5173 / GitHub Pages)", [
        "React Components (Navbar, EventCard, FilterPanel, Modals)",
        "React Router DOM with client-side protected route guards",
        "AuthContext global session management (User state, Token)",
        "Axios HTTP client with Bearer token request interceptor"
    ], BRAND_BLUE),
    ("Tier 2: Application Logic Tier", "Express.js RESTful API Server (Port 5000 / Render)", [
        "Security Pipeline: Helmet headers, CORS policies, Rate limiting",
        "Authentication & Authorization: JWT verification & RBAC check",
        "Input Validation: express-validator sanitization & bounds rules",
        "Resource Controllers: Auth, Events, Registrations, Admin"
    ], GREEN),
    ("Tier 3: Data Persistence Tier", "MongoDB Database Server (campus_events / Atlas)", [
        "Users Collection: Hashed passwords, Student/Admin roles",
        "Events Collection: Capacity, remainingSeats, indexes",
        "Registrations Collection: Unique compound index { student, event }",
        "Atomic Operators: findOneAndUpdate with $expr and $inc"
    ], INDIGO)
]

for i, (name, subtitle, points, color) in enumerate(tiers):
    top = Inches(1.5 + i * 1.8)
    add_card(s7, Inches(0.8), top, Inches(11.7), Inches(1.6))

    accent = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), top, Inches(0.15), Inches(1.6))
    accent.fill.solid()
    accent.fill.fore_color.rgb = color
    accent.line.fill.background()

    tb = s7.shapes.add_textbox(Inches(1.1), top + Inches(0.1), Inches(11.2), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"{name}  —  "
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = color
    r = p.add_run()
    r.text = subtitle
    r.font.size = Pt(11.5)
    r.font.color.rgb = TEXT_DARK

    p_list = tf.add_paragraph()
    p_list.space_before = Pt(4)
    p_list.text = "   • " + "   • ".join(points)
    p_list.font.size = Pt(10.5)
    p_list.font.color.rgb = TEXT_MUTED

set_speaker_notes(s7, "Slide 7 shows our 3-Tier Architecture. The client tier executes in the browser using React and Axios. Requests travel over HTTP to the Express backend tier, where security, authentication, and validation middleware process requests. Finally, controllers communicate with MongoDB using Mongoose schemas and atomic operators.")

# ==============================================================================
# SLIDE 8: DATABASE DESIGN & SCHEMA
# ==============================================================================
s8 = prs.slides.add_slide(blank_slide_layout)
add_header(s8, "Database Design: Mongoose Schemas & Indexing Strategy", "Data Models")

colls = [
    ("users Collection", "Student & Faculty Admin Accounts", [
        ("name", "String (2-60 chars, Required)"),
        ("email", "String (Unique index, normalized lowercase)"),
        ("password", "String (Bcrypt hashed, 10 salt rounds)"),
        ("role", "Enum: ['student', 'admin'] (Default: student)"),
        ("department", "String (College branch / department)"),
        ("studentId", "String (Campus Roll Number)"),
        ("timestamps", "Date (createdAt, updatedAt auto-generated)")
    ], BRAND_BLUE),
    ("events Collection", "Campus Symposiums, Hackathons, Fests", [
        ("title & desc", "String (Text search indexed, required)"),
        ("category", "Enum: Technical, Workshop, Cultural, Sports..."),
        ("eventDate", "Date (Compound index with status: 1)"),
        ("venue & time", "String (Auditorium/Lab, start & end time)"),
        ("capacity", "Number (Max seats limit, positive integer >= 1)"),
        ("registeredCount", "Number (Live enrolled attendee count)"),
        ("remainingSeats", "Virtual: Math.max(0, capacity - registered)")
    ], GREEN),
    ("registrations Collection", "Attendee Seat Reservations & Audit", [
        ("student", "ObjectId (ref: User, compound unique index)"),
        ("event", "ObjectId (ref: Event, compound unique index)"),
        ("status", "Enum: ['confirmed', 'cancelled']"),
        ("registeredAt", "Date (Seat booking timestamp)"),
        ("cancelledAt", "Date (Cancellation audit timestamp)"),
        ("notes", "String (Optional dietary or team requirements)"),
        ("UNIQUE KEY", "{ student: 1, event: 1 } (Prevents duplicates)")
    ], INDIGO)
]

for i, (cname, cdesc, fields, color) in enumerate(colls):
    left = Inches(0.8 + i * 3.95)
    top = Inches(1.5)
    add_card(s8, left, top, Inches(3.8), Inches(5.2))

    hdr = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(3.8), Inches(0.8))
    hdr.fill.solid()
    hdr.fill.fore_color.rgb = color
    hdr.line.fill.background()
    h_tf = hdr.text_frame
    h_tf.word_wrap = True
    p = h_tf.paragraphs[0]
    p.text = cname
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT
    p.alignment = PP_ALIGN.CENTER
    p2 = h_tf.add_paragraph()
    p2.text = cdesc
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = RGBColor(226, 232, 240)
    p2.alignment = PP_ALIGN.CENTER

    tb = s8.shapes.add_textbox(left + Inches(0.15), top + Inches(0.9), Inches(3.5), Inches(4.1))
    tf = tb.text_frame
    tf.word_wrap = True
    for fname, fdesc in fields:
        p = tf.add_paragraph()
        p.space_before = Pt(4)
        r1 = p.add_run()
        r1.text = f"• {fname}: "
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = TEXT_DARK
        r2 = p.add_run()
        r2.text = fdesc
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_MUTED

set_speaker_notes(s8, "Our MongoDB database is named campus_events and has three collections: users, events, and registrations. We use Mongoose virtuals for remainingSeats calculation, and compound unique indexes to guarantee that no student can hold multiple active registrations for the same event.")

# ==============================================================================
# SLIDE 9: FRONTEND DESIGN & USER INTERFACE
# ==============================================================================
s9 = prs.slides.add_slide(blank_slide_layout)
add_header(s9, "Frontend User Interface: Modern University Design System", "Visual Interface")

# Embed real screenshot 01 (Landing Page) & 02 (Events Catalog)
img1_path = os.path.join(SCREENSHOTS_DIR, "01_landing_page.png")
img2_path = os.path.join(SCREENSHOTS_DIR, "02_events_catalog.png")

if os.path.exists(img1_path):
    add_card(s9, Inches(0.8), Inches(1.5), Inches(5.7), Inches(4.3))
    s9.shapes.add_picture(img1_path, Inches(0.9), Inches(1.6), Inches(5.5), Inches(3.4))
    tb1 = s9.shapes.add_textbox(Inches(0.9), Inches(5.05), Inches(5.5), Inches(0.65))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Campus Landing Page (Hero, Live Stats, Category Cards)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE

if os.path.exists(img2_path):
    add_card(s9, Inches(6.8), Inches(1.5), Inches(5.7), Inches(4.3))
    s9.shapes.add_picture(img2_path, Inches(6.9), Inches(1.6), Inches(5.5), Inches(3.4))
    tb2 = s9.shapes.add_textbox(Inches(6.9), Inches(5.05), Inches(5.5), Inches(0.65))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Events Explorer (Keyword Search, Category Pills, Capacity Progress)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = GREEN

# Highlights Card at bottom
c_bot = add_card(s9, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.95))
tb_bot = s9.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.85))
tf_bot = tb_bot.text_frame
p = tf_bot.paragraphs[0]
p.text = "UI Design System Principles:  "
p.font.bold = True
p.font.size = Pt(11)
p.font.color.rgb = TEXT_DARK
r = p.add_run()
r.text = "Navy blue & indigo university theme, Bootstrap 5 responsive grid, real-time search filtering, color-coded capacity indicators, and clean accessible form inputs."
r.font.size = Pt(11)
r.font.color.rgb = TEXT_MUTED

set_speaker_notes(s9, "These are genuine screenshots captured directly from our running React application. On the left is the university landing page with hero banner and platform statistics. On the right is the events explorer showing category filter pills, keyword search, and dynamic seat capacity progress bars.")

# ==============================================================================
# SLIDE 10: STUDENT REGISTRATION WORKFLOW
# ==============================================================================
s10 = prs.slides.add_slide(blank_slide_layout)
add_header(s10, "Student Registration Workflow: Atomic Seat Booking Flow", "Transaction Logic")

img3_path = os.path.join(SCREENSHOTS_DIR, "03_event_details.png")
if os.path.exists(img3_path):
    add_card(s10, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.2))
    s10.shapes.add_picture(img3_path, Inches(0.9), Inches(1.6), Inches(5.5), Inches(4.4))
    tb = s10.shapes.add_textbox(Inches(0.9), Inches(6.1), Inches(5.5), Inches(0.5))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Live Event Details Page with Remaining Seats Gauge & Booking Action"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE

# Workflow Steps Card
c_wf = add_card(s10, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.2))
tb_wf = s10.shapes.add_textbox(Inches(7.1), Inches(1.7), Inches(5.1), Inches(4.8))
tf_wf = tb_wf.text_frame
tf_wf.word_wrap = True
p = tf_wf.paragraphs[0]
p.text = "Booking Lifecycle & Concurrency Guard"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = BRAND_BLUE

w_steps = [
    ("1. Student Authentication", "Student signs in with credentials; JWT Bearer token is stored in client session."),
    ("2. Event Selection & Review", "Student selects an event, reviews date, venue, syllabus, and available seats."),
    ("3. Deadline & Status Verification", "Server verifies event is 'published' and registration deadline has not elapsed."),
    ("4. Duplicate Booking Check", "Queries existing registration records. If active, rejects duplicate submission."),
    ("5. Atomic Seat Increment ($inc)", "MongoDB executes findOneAndUpdate with condition registeredCount < capacity. Prevents overselling completely."),
    ("6. Registration Record Created", "Document created in registrations collection. Student receives confirmed status badge.")
]

for title, desc in w_steps:
    p = tf_wf.add_paragraph()
    p.space_before = Pt(5)
    r1 = p.add_run()
    r1.text = f"{title}\n"
    r1.font.bold = True
    r1.font.size = Pt(11.5)
    r1.font.color.rgb = GREEN
    r2 = p.add_run()
    r2.text = desc
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = TEXT_MUTED

set_speaker_notes(s10, "Slide 10 details our atomic registration workflow. When a student clicks Register, our backend executes a conditional update checking that current count is strictly less than capacity. This guarantees no two students can take the same remaining seat. On the left is the live event details page.")

# ==============================================================================
# SLIDE 11: ADMINISTRATOR DASHBOARD
# ==============================================================================
s11 = prs.slides.add_slide(blank_slide_layout)
add_header(s11, "Administrator Console: Live Management & Analytics", "Administrative Panel")

img7_path = os.path.join(SCREENSHOTS_DIR, "07_admin_dashboard.png")
img8_path = os.path.join(SCREENSHOTS_DIR, "08_create_event.png")

if os.path.exists(img7_path):
    add_card(s11, Inches(0.8), Inches(1.5), Inches(5.7), Inches(4.3))
    s11.shapes.add_picture(img7_path, Inches(0.9), Inches(1.6), Inches(5.5), Inches(3.4))
    tb1 = s11.shapes.add_textbox(Inches(0.9), Inches(5.05), Inches(5.5), Inches(0.65))
    tf1 = tb1.text_frame
    p = tf1.paragraphs[0]
    p.text = "Admin Dashboard (KPI Metrics, Capacity Progress Bars, Actions)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE

if os.path.exists(img8_path):
    add_card(s11, Inches(6.8), Inches(1.5), Inches(5.7), Inches(4.3))
    s11.shapes.add_picture(img8_path, Inches(6.9), Inches(1.6), Inches(5.5), Inches(3.4))
    tb2 = s11.shapes.add_textbox(Inches(6.9), Inches(5.05), Inches(5.5), Inches(0.65))
    tf2 = tb2.text_frame
    p = tf2.paragraphs[0]
    p.text = "Event Publication Form (Date Chronology, Capacity & Status Rules)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = GREEN

c_bot = add_card(s11, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.95))
tb_bot = s11.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.85))
tf_bot = tb_bot.text_frame
p = tf_bot.paragraphs[0]
p.text = "Faculty Administrative Powers:  "
p.font.bold = True
p.font.size = Pt(11)
p.font.color.rgb = TEXT_DARK
r = p.add_run()
r.text = "Live analytics cards (Total Events, Students, Confirmed Bookings), event CRUD with capacity safety bounds, attendee search by roll ID, and automated CSV spreadsheet export."
r.font.size = Pt(11)
r.font.color.rgb = TEXT_MUTED

set_speaker_notes(s11, "Slide 11 shows actual screenshots of our Administrator Console. The left image demonstrates the KPI dashboard displaying total events, student accounts, and confirmed bookings. The right image shows the event creation form with validation rules enforcing that registration deadlines precede event dates.")

# ==============================================================================
# SLIDE 12: BACKEND & REST API INTEGRATION
# ==============================================================================
s12 = prs.slides.add_slide(blank_slide_layout)
add_header(s12, "Backend & API Integration: RESTful Endpoints Architecture", "Server Architecture")

endpoints = [
    ("POST /api/auth/register", "Public", "Student registration with bcrypt password hashing and JWT issuance."),
    ("POST /api/auth/login", "Public", "User authentication; returns token and role payload."),
    ("GET  /api/events", "Public", "Query events with search, category filters, and pagination parameters."),
    ("GET  /api/events/:id", "Public / Auth", "Event details with current user's registration status badge."),
    ("POST /api/events", "Admin Only", "Create campus event with capacity and chronological date rules."),
    ("POST /api/registrations/:id", "Student", "Atomic seat reservation with conditional concurrency check."),
    ("DELETE /api/registrations/:id", "Student", "Cancel upcoming registration; restores seat capacity count."),
    ("GET  /api/admin/dashboard", "Admin Only", "Platform metrics, category distribution, and capacity table."),
    ("GET  /api/admin/registrations/export/:id", "Admin Only", "Generates downloadable CSV file stream of event attendees.")
]

for i, (route, access, desc) in enumerate(endpoints):
    col = i % 2
    row = i // 2
    left = Inches(0.8 + col * 6.0)
    top = Inches(1.4 + row * 1.1)
    add_card(s12, left, top, Inches(5.7), Inches(0.95))

    tb = s12.shapes.add_textbox(left + Inches(0.15), top + Inches(0.05), Inches(5.4), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = route
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE
    
    r_badge = p.add_run()
    r_badge.text = f"  [{access}]"
    r_badge.font.size = Pt(9.5)
    r_badge.font.bold = True
    r_badge.font.color.rgb = GREEN if "Public" in access else AMBER

    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = TEXT_MUTED

set_speaker_notes(s12, "Here is our RESTful API endpoint architecture. Routes are organized into four resources: authentication, events, registrations, and administrative reporting. All private routes are guarded by JWT verification and role authorization middleware.")

# ==============================================================================
# SLIDE 13: INPUT VALIDATION & SECURITY SAFEGUARDS
# ==============================================================================
s13 = prs.slides.add_slide(blank_slide_layout)
add_header(s13, "Input Validation & Security Safeguards", "System Security")

sec_cards = [
    ("Client-Side Validation (React)", [
        "Email regex pattern validation",
        "Minimum password length (6+ characters)",
        "Confirm password matching parity",
        "Required field indicators and instant feedback"
    ], BRAND_BLUE),
    ("Server-Side Validation (Express)", [
        "express-validator schema sanitization",
        "Chronological checks: deadline <= event date",
        "Capacity bounds: positive integer >= 1",
        "ObjectId format verification (isMongoId)"
    ], GREEN),
    ("Authentication & Authorization", [
        "Bcrypt password hashing (10 salt rounds)",
        "Stateless 7-day HMAC-SHA256 JWT tokens",
        "Role-Based Access Control (Student vs Admin)",
        "ProtectedRoute and AdminRoute guards"
    ], INDIGO),
    ("Infrastructure Hardening", [
        "Helmet HTTP security response headers",
        "CORS restricted to authorized origins",
        "express-rate-limit brute force protection",
        "Environment variables (.env) for credentials"
    ], AMBER)
]

for i, (title, points, color) in enumerate(sec_cards):
    col = i % 2
    row = i // 2
    left = Inches(0.8 + col * 6.0)
    top = Inches(1.5 + row * 2.7)
    add_card(s13, left, top, Inches(5.7), Inches(2.4))

    accent = s13.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.12), Inches(2.4))
    accent.fill.solid()
    accent.fill.fore_color.rgb = color
    accent.line.fill.background()

    tb = s13.shapes.add_textbox(left + Inches(0.3), top + Inches(0.15), Inches(5.2), Inches(2.1))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = color

    for pt in points:
        p2 = tf.add_paragraph()
        p2.space_before = Pt(3)
        p2.text = f"✔  {pt}"
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_DARK

set_speaker_notes(s13, "Slide 13 outlines our security architecture. We never rely solely on frontend validation. Every request is re-validated on the server with express-validator. Passwords use bcrypt with 10 salt rounds, sessions use stateless JWTs, and Helmet secures HTTP headers.")

# ==============================================================================
# SLIDE 14: MONGODB DEMONSTRATION & DATA INTEGRITY
# ==============================================================================
s14 = prs.slides.add_slide(blank_slide_layout)
add_header(s14, "MongoDB Data Integrity: Event & Registration Persistence", "Database Verification")

img9_path = os.path.join(SCREENSHOTS_DIR, "09_attendee_roster.png")
if os.path.exists(img9_path):
    add_card(s14, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.2))
    s14.shapes.add_picture(img9_path, Inches(0.9), Inches(1.6), Inches(5.5), Inches(4.3))
    tb = s14.shapes.add_textbox(Inches(0.9), Inches(6.05), Inches(5.5), Inches(0.55))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Attendee Roster View Generated from MongoDB Registrations Collection"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE

# Document Structure Card
c_doc = add_card(s14, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.2))
tb_doc = s14.shapes.add_textbox(Inches(7.1), Inches(1.7), Inches(5.1), Inches(4.8))
tf_doc = tb_doc.text_frame
tf_doc.word_wrap = True
p = tf_doc.paragraphs[0]
p.text = "Physical Document Separation"
p.font.size = Pt(17)
p.font.bold = True
p.font.color.rgb = GREEN

doc_text = [
    ("Event Document (events collection):", "Stores title, category, date, venue, capacity (e.g. 120), and registeredCount (e.g. 3)."),
    ("Registration Document (registrations):", "Stores individual attendee records referencing student: ObjectId and event: ObjectId, with status: 'confirmed'."),
    ("Referential Integrity via Mongoose:", "populate('student') and populate('event') dynamically join records without monolithic tables."),
    ("Compound Unique Index Integrity:", "index({ student: 1, event: 1 }, { unique: true }) ensures that even if concurrent network calls race, only ONE registration document is ever created.")
]

for t, d in doc_text:
    p = tf_doc.add_paragraph()
    p.space_before = Pt(6)
    r1 = p.add_run()
    r1.text = f"{t}\n"
    r1.font.bold = True
    r1.font.size = Pt(11.5)
    r1.font.color.rgb = BRAND_BLUE
    r2 = p.add_run()
    r2.text = d
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = TEXT_MUTED

set_speaker_notes(s14, "In Slide 14, we demonstrate data persistence in MongoDB. Event records and registration records are stored separately for referential integrity. On the left is the live attendee roster. The compound unique index guarantees that registeredCount and actual registration documents remain 100% synchronized.")

# ==============================================================================
# SLIDE 15: TESTING & RESULTS
# ==============================================================================
s15 = prs.slides.add_slide(blank_slide_layout)
add_header(s15, "Testing Methodology & Verification Results", "Quality Assurance")

# Metrics Cards Top
kpis = [
    ("21 / 21", "Automated API Tests Passed", GREEN),
    ("25 / 25", "Documented Test Scenarios Passed", BRAND_BLUE),
    ("100%", "Test Pass Rate", INDIGO),
    ("1.16s", "Vite Production Build Time", AMBER)
]

for i, (stat, label, color) in enumerate(kpis):
    left = Inches(0.8 + i * 2.95)
    top = Inches(1.4)
    add_card(s15, left, top, Inches(2.8), Inches(1.15))
    tb = s15.shapes.add_textbox(left + Inches(0.1), top + Inches(0.1), Inches(2.6), Inches(0.95))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = stat
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = color
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = label
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = TEXT_MUTED
    p2.alignment = PP_ALIGN.CENTER

# Test Summary Table Card
c_tbl = add_card(s15, Inches(0.8), Inches(2.75), Inches(11.7), Inches(4.0))
tb_tbl = s15.shapes.add_textbox(Inches(1.0), Inches(2.9), Inches(11.3), Inches(3.7))
tf_tbl = tb_tbl.text_frame
tf_tbl.word_wrap = True

p = tf_tbl.paragraphs[0]
p.text = "Key Verified Test Scenarios (Sample Matrix):"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = TEXT_DARK

test_rows = [
    ("TC-01", "Student Registration", "Valid name, email, and password payload", "201 Created; JWT token & profile returned", "PASS"),
    ("TC-02", "Duplicate Email Rejection", "Register with already existing email", "400 Bad Request: Account already exists", "PASS"),
    ("TC-08", "Role Authorization Guard", "Student JWT accesses /api/admin/dashboard", "403 Forbidden: User role 'student' unauthorized", "PASS"),
    ("TC-12", "Atomic Seat Reservation", "Student registers for open published event", "201 Created; registeredCount atomically +1", "PASS"),
    ("TC-13", "Duplicate Event Booking", "Same student attempts second registration", "400 Bad Request: Already registered for event", "PASS"),
    ("TC-17", "Cancellation & Seat Return", "Student cancels seat for upcoming event", "200 OK; registeredCount atomically -1", "PASS"),
    ("TC-21", "CSV Roster Streaming", "Admin requests /registrations/export/:id", "200 OK; text/csv file stream generated", "PASS")
]

for tid, tname, tinput, texp, tstat in test_rows:
    p = tf_tbl.add_paragraph()
    p.space_before = Pt(3)
    r1 = p.add_run()
    r1.text = f"[{tid}] {tname}: "
    r1.font.bold = True
    r1.font.size = Pt(10)
    r1.font.color.rgb = BRAND_BLUE
    r2 = p.add_run()
    r2.text = f"{texp}  —  "
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = TEXT_MUTED
    r3 = p.add_run()
    r3.text = f"[{tstat}]"
    r3.font.bold = True
    r3.font.size = Pt(10)
    r3.font.color.rgb = GREEN

set_speaker_notes(s15, "Slide 15 shows our rigorous testing methodology. We executed 21 automated integration tests via backend/test/api.test.js with a 100% pass rate. We verified duplicate rejection, role authorization blocking, atomic increment, cancellation decrement, and CSV downloads.")

# ==============================================================================
# SLIDE 16: GITHUB REPOSITORY & DEPLOYMENT ARCHITECTURE
# ==============================================================================
s16 = prs.slides.add_slide(blank_slide_layout)
add_header(s16, "GitHub Repository & Deployment Architecture", "DevOps & Hosting")

c_repo = add_card(s16, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.2))
tb_repo = s16.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.1), Inches(4.8))
tf_repo = tb_repo.text_frame
tf_repo.word_wrap = True
p = tf_repo.paragraphs[0]
p.text = "GitHub Version Control"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = BRAND_BLUE

repo_items = [
    ("Repository URL", "https://github.com/Tamil-k-iniyan/Campus-Event-Registration-"),
    ("Branch & Commits", "main branch with descriptive commits, structured .gitignore excluding secrets, and MIT License."),
    ("CI/CD Automation", "GitHub Actions workflow (.github/workflows/deploy.yml) automatically builds Vite frontend and publishes to GitHub Pages."),
    ("Academic Documentation", "Repository includes PROJECT_REPORT.docx, DATABASE_SCHEMA.md, API_DOCUMENTATION.md, and TEST_CASES.md.")
]
for title, desc in repo_items:
    p = tf_repo.add_paragraph()
    p.space_before = Pt(8)
    r1 = p.add_run()
    r1.text = f"• {title}:\n"
    r1.font.bold = True
    r1.font.size = Pt(12)
    r1.font.color.rgb = TEXT_DARK
    r2 = p.add_run()
    r2.text = desc
    r2.font.size = Pt(11)
    r2.font.color.rgb = BRAND_BLUE if "http" in desc else TEXT_MUTED

c_dep = add_card(s16, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.2))
tb_dep = s16.shapes.add_textbox(Inches(7.1), Inches(1.7), Inches(5.1), Inches(4.8))
tf_dep = tb_dep.text_frame
tf_dep.word_wrap = True
p = tf_dep.paragraphs[0]
p.text = "Deployment & Hosting Options"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = GREEN

dep_items = [
    ("Live Frontend (GitHub Pages)", "https://tamil-k-iniyan.github.io/Campus-Event-Registration-/"),
    ("Cloud API Hosting (Render)", "Ready-to-deploy render.yaml blueprint configured for free Node.js service hosting."),
    ("Cloud Database (MongoDB Atlas)", "Free MongoDB M0 Cluster connection string configured via MONGO_URI."),
    ("Local Execution Commands", "Backend: cd backend && npm run dev (Port 5000)\nFrontend: cd frontend && npm run dev (Port 5173)")
]
for title, desc in dep_items:
    p = tf_dep.add_paragraph()
    p.space_before = Pt(8)
    r1 = p.add_run()
    r1.text = f"• {title}:\n"
    r1.font.bold = True
    r1.font.size = Pt(12)
    r1.font.color.rgb = TEXT_DARK
    r2 = p.add_run()
    r2.text = desc
    r2.font.size = Pt(11)
    r2.font.color.rgb = BRAND_BLUE if "http" in desc else TEXT_MUTED

set_speaker_notes(s16, "Slide 16 details version control and deployment. Our GitHub repository is Tamil-k-iniyan/Campus-Event-Registration-. The frontend deploys via GitHub Actions, and the repository includes a render.yaml blueprint for deploying the Express API to Render with MongoDB Atlas.")

# ==============================================================================
# SLIDE 17: FUTURE ENHANCEMENTS
# ==============================================================================
s17 = prs.slides.add_slide(blank_slide_layout)
add_header(s17, "Future Enhancements: Roadmap for Production Scaling", "Future Work")

futures = [
    ("1. Contactless QR-Code Ticketing", "Generate signed QR code entry passes for confirmed students; provide a volunteer mobile scanner app for rapid venue admission."),
    ("2. Automated Waitlist Queue", "When capacity hits maximum, allow students to join an automated queue that promotes attendees when cancellations occur."),
    ("3. Automated Certificate Generation", "Dynamically generate verifiable PDF participation and winner certificates stamped with unique verification hashes."),
    ("4. Push & Email Notifications", "Integrate Nodemailer and web push notifications to alert students of registration confirmations and 24-hour event reminders."),
    ("5. Academic Calendar Integration", "Provide 1-click 'Add to Google Calendar' and Apple iCal (.ics) exports with automated schedule synchronization."),
    ("6. Advanced Participation Analytics", "Provide departmental Dean reports tracking student engagement trends across academic semesters.")
]

for i, (title, desc) in enumerate(futures):
    col = i % 2
    row = i // 2
    left = Inches(0.8 + col * 6.0)
    top = Inches(1.5 + row * 1.8)
    add_card(s17, left, top, Inches(5.7), Inches(1.6))

    tb = s17.shapes.add_textbox(left + Inches(0.3), top + Inches(0.15), Inches(5.2), Inches(1.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE
    
    p2 = tf.add_paragraph()
    p2.space_before = Pt(4)
    p2.text = desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_MUTED

set_speaker_notes(s17, "Looking forward, we have identified high-value enhancements for campus deployment: contactless QR-code ticket scanning at venue doors, automated waitlist queues, automatic PDF participation certificates, and calendar synchronization.")

# ==============================================================================
# SLIDE 18: CONCLUSION & VIVA VOCE
# ==============================================================================
s18 = prs.slides.add_slide(blank_slide_layout)
# Dark navy finish
s18_bg = s18.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
s18_bg.fill.solid()
s18_bg.fill.fore_color.rgb = NAVY
s18_bg.line.fill.background()

c_end = add_card(s18, Inches(1.5), Inches(1.2), Inches(10.333), Inches(5.1), bg_color=INDIGO, border_color=BRAND_BLUE)
tb_end = s18.shapes.add_textbox(Inches(1.8), Inches(1.4), Inches(9.7), Inches(4.7))
tf_end = tb_end.text_frame
tf_end.word_wrap = True

p = tf_end.paragraphs[0]
p.text = "Project Conclusion"
p.font.size = Pt(24)
p.font.bold = True
p.font.color.rgb = RGBColor(56, 189, 248)

concl_points = [
    "Successfully delivered a complete, professional, full-stack campus event management web application.",
    "Eliminated event overbooking through atomic concurrency-safe seat reservation algorithms in MongoDB.",
    "Prevented duplicate registrations using database-level unique compound indexing ({ student: 1, event: 1 }).",
    "Empowered faculty with live analytics dashboards, capacity utilization progress bars, and instant CSV exports.",
    "Validated 100% of API endpoints through automated integration testing and rigorous input sanitization."
]
for pt in concl_points:
    p2 = tf_end.add_paragraph()
    p2.space_before = Pt(6)
    p2.text = f"✔  {pt}"
    p2.font.size = Pt(12.5)
    p2.font.color.rgb = TEXT_LIGHT

p_ty = tf_end.add_paragraph()
p_ty.space_before = Pt(18)
p_ty.text = "THANK YOU!"
p_ty.font.size = Pt(28)
p_ty.font.bold = True
p_ty.font.color.rgb = TEXT_LIGHT
p_ty.alignment = PP_ALIGN.CENTER

p_q = tf_end.add_paragraph()
p_q.text = "Questions & Answers  |  Tamilzhniyan K.  |  URK24CS1140"
p_q.font.size = Pt(13)
p_q.font.bold = True
p_q.font.color.rgb = RGBColor(148, 163, 184)
p_q.alignment = PP_ALIGN.CENTER

set_speaker_notes(s18, "In conclusion, the Campus Event Registration and Management System successfully demonstrates how full-stack web engineering solves real campus logistical problems. Thank you for your time and guidance. I welcome your questions.")

# Save presentation
out_path = os.path.abspath(os.path.join(os.path.dirname(__file__), '../../Campus_Event_Registration_Management_System.pptx'))
prs.save(out_path)
print(f"[PPTX Generator] Successfully generated 18 slides to: {out_path}")
